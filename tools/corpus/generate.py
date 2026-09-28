#!/usr/bin/env python3
"""Deterministic generator for the Waffle test-file corpus.

Usage (from repo root):
    python3 -m venv "$VENV" && "$VENV/bin/pip" install -r testdata/gen/requirements.txt
    "$VENV/bin/python" testdata/gen/generate.py            # everything
    "$VENV/bin/python" testdata/gen/generate.py --skip-large

Every output is byte-for-byte reproducible for the pinned library versions in
requirements.txt: RNGs use fixed seeds, and every zip container is rewritten with
fixed timestamps and fixed docProps dates (see normalize_zip).
"""
import argparse
import datetime as dt
import io
import os
import random
import re
import shutil
import struct
import sys
import time
import zipfile
import zlib

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "..", "..", "testdata", "files")
VENDOR = os.path.join(HERE, "vendor")

FIXED_DT = dt.datetime(2024, 1, 1, 0, 0, 0)
FIXED_ZIP_DT = (1980, 1, 1, 0, 0, 0)


def out(name):
    return os.path.join(OUT, name)


# ---------------------------------------------------------------------------
# zip helpers
# ---------------------------------------------------------------------------

_DCTERMS_RE = re.compile(rb"(<dcterms:(created|modified)[^>]*>)[^<]*(</dcterms:\2>)")
_ODF_DATE_RE = re.compile(rb"(<meta:(creation-date|date)>)[^<]*(</meta:\2>)|(<dc:date>)[^<]*(</dc:date>)")


def normalize_zip(path, transform=None):
    """Rewrite a zip with fixed timestamps/attrs, preserving entry order and
    compression method. Pins docProps created/modified dates. `transform` may
    rewrite individual entries: transform(name, bytes) -> bytes."""
    with zipfile.ZipFile(path) as zin:
        entries = [(i, zin.read(i.filename)) for i in zin.infolist()]
    tmp = path + ".tmp"
    with zipfile.ZipFile(tmp, "w") as zout:
        for info, data in entries:
            name = info.filename
            if name in ("docProps/core.xml",):
                data = _DCTERMS_RE.sub(rb"\g<1>2024-01-01T00:00:00Z\3", data)
            if name == "meta.xml":
                data = re.sub(rb"<meta:creation-date>[^<]*</meta:creation-date>",
                              b"<meta:creation-date>2024-01-01T00:00:00</meta:creation-date>", data)
                data = re.sub(rb"<dc:date>[^<]*</dc:date>", b"<dc:date>2024-01-01T00:00:00</dc:date>", data)
            if transform:
                data = transform(name, data)
            zi = zipfile.ZipInfo(name, date_time=FIXED_ZIP_DT)
            zi.compress_type = info.compress_type
            zi.external_attr = 0o644 << 16
            zi.create_system = 0
            zout.writestr(zi, data)
    os.replace(tmp, path)


def tiny_png(w=24, h=16):
    """Deterministic RGB PNG: left half red, right half blue."""
    raw = b""
    for y in range(h):
        raw += b"\x00" + b"".join(
            (b"\xd0\x20\x20" if x < w // 2 else b"\x20\x40\xd0") for x in range(w))

    def chunk(t, d):
        c = struct.pack(">I", len(d)) + t + d
        return c + struct.pack(">I", zlib.crc32(t + d) & 0xFFFFFFFF)

    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw, 9)) + chunk(b"IEND", b""))


def col_letter(idx):  # 0-based
    s = ""
    idx += 1
    while idx:
        idx, r = divmod(idx - 1, 26)
        s = chr(65 + r) + s
    return s


# ---------------------------------------------------------------------------
# 1. styled.xlsx (openpyxl)
# ---------------------------------------------------------------------------

def gen_styled():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Color, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Styled"
    wb.properties.creator = "xcell-corpus"
    wb.properties.created = FIXED_DT
    wb.properties.modified = FIXED_DT

    hdr_font = Font(bold=True, color="FFFFFFFF")
    hdr_fill = PatternFill("solid", fgColor="FF305496")
    for c, text in zip("ABCD", ["Feature", "Sample", "Expected display", "Merges"]):
        cell = ws[f"{c}1"]
        cell.value = text
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    def row(r, label, value, expected=None, **style):
        ws[f"A{r}"] = label
        cell = ws[f"B{r}"]
        cell.value = value
        for k, v in style.items():
            setattr(cell, k, v)
        if expected is not None:
            ws[f"C{r}"] = expected
            ws[f"C{r}"].number_format = "@"

    thin, med, thick = Side("thin"), Side("medium"), Side("thick")
    dashed_red = Side("dashed", color="FFFF0000")
    # --- fonts / fills / borders / alignment
    row(2, "Bold", "Bold text", font=Font(bold=True))
    row(3, "Italic", "Italic text", font=Font(italic=True))
    row(4, "Underline", "Underlined text", font=Font(underline="single"))
    row(5, "Strikethrough", "Struck text", font=Font(strike=True))
    row(6, "Georgia 18", "Big Georgia", font=Font(name="Georgia", size=18))
    row(7, "Courier New 8", "Small Courier", font=Font(name="Courier New", size=8))
    row(8, "Font color rgb", "Red FF0000", font=Font(color="FFFF0000"))
    row(9, "Font color theme", "Theme 4 tint -0.25", font=Font(color=Color(theme=4, tint=-0.25)))
    row(10, "Font color indexed", "Indexed 12", font=Font(color=Color(indexed=12)))
    row(11, "Solid fill", "Yellow fill", fill=PatternFill("solid", fgColor="FFFFFF00"))
    row(12, "Border thin", "thin all", border=Border(left=thin, right=thin, top=thin, bottom=thin))
    row(13, "Border medium", "medium all", border=Border(left=med, right=med, top=med, bottom=med))
    row(14, "Border thick", "thick all", border=Border(left=thick, right=thick, top=thick, bottom=thick))
    row(15, "Border dashed red", "dashed red", border=Border(left=dashed_red, right=dashed_red,
                                                               top=dashed_red, bottom=dashed_red))
    row(16, "Align right/top", "right-top", alignment=Alignment(horizontal="right", vertical="top"))
    row(17, "Align center/center", "centered", alignment=Alignment(horizontal="center", vertical="center"))
    row(18, "Wrap text", "This is a long piece of text that should wrap inside the cell",
        alignment=Alignment(wrap_text=True, vertical="top"))
    row(19, "Indent 2", "indented", alignment=Alignment(horizontal="left", indent=2))

    # --- number formats: (label, value, format, expected display in en-US)
    nf = [
        ("0.00", 1234.5, "0.00", "1234.50"),
        ("#,##0.00", 1234567.891, "#,##0.00", "1,234,567.89"),
        ("0%", 0.256, "0%", "26%"),
        ("0.00%", 0.256, "0.00%", "25.60%"),
        ("m/d/yyyy", dt.datetime(2024, 3, 5), "m/d/yyyy", "3/5/2024"),
        ("dd-mmm-yyyy", dt.datetime(2024, 3, 5), "dd-mmm-yyyy", "05-Mar-2024"),
        ("yyyy-mm-dd hh:mm", dt.datetime(2024, 3, 5, 14, 30), "yyyy-mm-dd hh:mm", "2024-03-05 14:30"),
        ("h:mm:ss AM/PM", dt.time(14, 30, 15), "h:mm:ss AM/PM", "2:30:15 PM"),
        ("hh:mm", dt.time(9, 5), "hh:mm", "09:05"),
        ("$ currency", 1234.5, '"$"#,##0.00', "$1,234.50"),
        ("INR currency", 123456.78, "[$₹-4009] #,##0.00", "₹ 123,456.78"),
        ("Accounting", -1234.5, '_("$"* #,##0.00_);_("$"* (#,##0.00);_("$"* "-"??_);_(@_)',
         " $ (1,234.50)"),
        ("[Red] negative", -42.5, "#,##0.00;[Red]-#,##0.00", "-42.50 (red)"),
        ("3-section zero", 0, '#,##0.00;[Red](#,##0.00);"zero"', "zero"),
        ("Text @", "00123", "@", "00123"),
        ("Scientific", 123456, "0.000E+00", "1.235E+05"),
        ("General", 0.1 + 0.2, "General", "0.3"),
    ]
    for i, (label, value, fmt, exp) in enumerate(nf):
        r = 20 + i  # rows 20..36
        row(r, label, value, exp)
        ws[f"B{r}"].number_format = fmt

    # --- hidden row / column
    ws["A37"] = "hidden row"
    ws["B37"] = 37
    ws.row_dimensions[37].hidden = True
    ws["A38"] = "after hidden row"
    ws["G1"] = "hidden column"
    ws["G2"] = 7
    ws.column_dimensions["G"].hidden = True
    ws["H1"] = "after hidden column"

    # --- merges
    ws["D2"] = "Merged across D2:F2"
    ws.merge_cells("D2:F2")
    ws["D4"] = "Merged down D4:D7"
    ws["D4"].alignment = Alignment(vertical="center")
    ws.merge_cells("D4:D7")
    ws["D9"] = "Merged block D9:F11"
    ws["D9"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.merge_cells("D9:F11")

    # --- sizes, freeze
    for c, w in {"A": 22, "B": 26, "C": 24, "D": 18, "E": 10, "F": 10, "H": 20}.items():
        ws.column_dimensions[c].width = w
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[16].height = 40
    ws.row_dimensions[18].height = 45
    ws.freeze_panes = "B2"

    path = out("styled.xlsx")
    wb.save(path)
    normalize_zip(path)


# ---------------------------------------------------------------------------
# 2. formulas.xlsx (xlsxwriter + shared-formula post-processing)
# ---------------------------------------------------------------------------

_CELL_F_RE = re.compile(r'(<c r="([A-Z]+)(\d+)"[^>]*>)<f>([^<]*)</f>')


def _make_shared(xml, cells, si, expected_formula):
    """Convert ordinary formulas in `cells` (list of refs, first = master) into a
    shared formula group. expected_formula(ref) gives the text each cell must
    currently contain (sanity check that the group really is shift-equivalent)."""
    ref = f"{cells[0]}:{cells[-1]}"
    want = set(cells)
    seen = {}

    def repl(m):
        cref = m.group(2) + m.group(3)
        if cref not in want:
            return m.group(0)
        f = m.group(4)
        exp = expected_formula(cref)
        if f != exp:
            raise AssertionError(f"{cref}: {f!r} != {exp!r}")
        seen[cref] = True
        if cref == cells[0]:
            return f'{m.group(1)}<f t="shared" ref="{ref}" si="{si}">{f}</f>'
        return f'{m.group(1)}<f t="shared" si="{si}"/>'

    xml = _CELL_F_RE.sub(repl, xml)
    assert len(seen) == len(cells), (cells, seen)
    return xml


def gen_formulas():
    import xlsxwriter

    path = out("formulas.xlsx")
    wb = xlsxwriter.Workbook(path)
    wb.set_properties({"author": "xcell-corpus", "created": FIXED_DT})
    calc = wb.add_worksheet("Calc")
    other = wb.add_worksheet("Other Sheet")
    bold = wb.add_format({"bold": True})

    other.write("A1", 100)
    other.write("B1", 0.18)
    other.write("A2", "text on other sheet")
    other.write("B2", 7)
    wb.define_name("TaxRate", "='Other Sheet'!$B$1")

    for c, h in zip("ABCDE", ["n", "double (shared)", "times F1 (shared, abs)", "label", "cross-sheet"]):
        calc.write(f"{c}1", h, bold)
    calc.write("F1", 3)
    calc.write("G1", "<- multiplier", bold)
    for i in range(10):
        r = i + 2
        n = i + 1
        calc.write_number(f"A{r}", n)
        calc.write_formula(f"B{r}", f"=A{r}*2", None, n * 2)
        calc.write_formula(f"C{r}", f"=A{r}*$F$1", None, n * 3)
        calc.write_formula(f"D{r}", f'=IF(A{r}>5,"big","small")', None, "big" if n > 5 else "small")
    calc.write_formula("E2", "='Other Sheet'!A1", None, 100)
    calc.write_formula("E3", "='Other Sheet'!B2*2", None, 14)
    calc.write_formula("E4", "=Calc!A2+'Other Sheet'!$A$1", None, 101)

    calc.write("A13", "totals", bold)
    calc.write_formula("B13", "=SUM(B2:B11)", None, 110)
    calc.write_formula("C13", "=SUM(C2:C11)", None, 165)

    calc.write("H1", "misc", bold)
    calc.write_formula("H2", "=SUM(A:A)", None, 55)
    calc.write_formula("H3", "=SUM(B:B)", None, 110 + 110)  # B13 total is included
    calc.write_formula("H4", "=TaxRate*100", None, 18)
    calc.write_formula("H5", "=SUM($A$2:$A$11)", None, 55)
    calc.write_formula("H6", "=A2/0", None, "#DIV/0!")
    calc.write_formula("H7", "=A2>0", None, True)

    calc.write("J1", "array", bold)
    calc.write_array_formula("J2:J4", "{=A2:A4*B2:B4}", None, 2)
    calc.write_number("J3", 8)
    calc.write_number("J4", 18)
    calc.write_array_formula("J6", "{=SUM(A2:A11*B2:B11)}", None, 770)
    wb.close()

    def transform(name, data):
        if name != "xl/worksheets/sheet1.xml":
            return data
        xml = data.decode("utf-8")
        xml = _make_shared(xml, [f"B{r}" for r in range(2, 12)], 0, lambda c: f"A{c[1:]}*2")
        xml = _make_shared(xml, [f"C{r}" for r in range(2, 12)], 1, lambda c: f"A{c[1:]}*$F$1")
        xml = _make_shared(xml, ["B13", "C13"], 2, lambda c: f"SUM({c[0]}2:{c[0]}11)")
        return xml.encode("utf-8")

    normalize_zip(path, transform)


# ---------------------------------------------------------------------------
# 3. rich.xlsx (xlsxwriter)
# ---------------------------------------------------------------------------

REGIONS = ["North", "South", "East", "West"]
PRODUCTS = ["Widget", "Gadget", "Doohickey", "Gizmo", "Thingamajig"]


def rich_rows():
    rng = random.Random(7)
    rows = []
    for i in range(20):
        units = rng.randint(1, 99)
        price = round(rng.uniform(1, 50), 2)
        rows.append([REGIONS[i % 4], PRODUCTS[i % 5], units, price, round(units * price, 2)])
    return rows


def gen_rich():
    import xlsxwriter

    path = out("rich.xlsx")
    wb = xlsxwriter.Workbook(path)
    wb.set_properties({"author": "xcell-corpus", "created": FIXED_DT, "title": "Rich corpus file"})
    summary = wb.add_worksheet("Summary")
    data = wb.add_worksheet("Data")
    hidden = wb.add_worksheet("Hidden Calc")
    notes = wb.add_worksheet("Notes")
    hidden.hide()

    bold = wb.add_format({"bold": True})
    bold_red = wb.add_format({"bold": True, "font_color": "#C00000"})
    ital_blue = wb.add_format({"italic": True, "font_color": "#0070C0"})
    red_fill = wb.add_format({"bg_color": "#FFC7CE", "font_color": "#9C0006"})
    money = wb.add_format({"num_format": "#,##0.00"})

    # Hidden sheet: list source + defined name
    for i, r in enumerate(REGIONS):
        hidden.write(i, 0, r)
    hidden.write("C1", "secret")
    wb.define_name("RegionList", "='Hidden Calc'!$A$1:$A$4")

    # Summary
    summary.write_rich_string("A1", bold_red, "Bold red", " plain ", ital_blue, "italic blue")
    summary.write("A3", "Region:", bold)
    summary.write("B3", "North")
    summary.data_validation("B3", {"validate": "list", "source": "=RegionList"})
    summary.write_comment("B3", "Pick a region from the list", {"author": "Xcell QA"})
    summary.write("A4", "Size:", bold)
    summary.write("B4", "M")
    summary.data_validation("B4", {"validate": "list", "source": ["S", "M", "L", "XL"]})
    summary.write_url("A6", "https://example.com/xcell", string="External link")
    summary.write_url("A7", "internal:'Data'!A1", string="Go to Data")
    summary.write_url("A8", "mailto:qa@example.com", string="Mail QA")
    summary.write_rich_string("A10", "Mixed: ", bold, "bold", " and ", ital_blue, "blue italic", " end")
    summary.set_column("A:A", 18)

    chart = wb.add_chart({"type": "column"})
    chart.add_series({"name": "Units", "categories": "='Data'!$B$2:$B$11",
                      "values": "='Data'!$C$2:$C$11"})
    chart.set_title({"name": "Units by product"})
    summary.insert_chart("D2", chart)

    # Data: table + conditional formats
    rows = rich_rows()
    data.add_table(0, 0, len(rows), 4, {
        "name": "SalesTable",
        "style": "Table Style Medium 2",
        "data": rows,
        "columns": [{"header": "Region"}, {"header": "Product"}, {"header": "Units"},
                    {"header": "Price", "format": money}, {"header": "Revenue", "format": money}],
    })
    data.conditional_format("C2:C21", {"type": "cell", "criteria": ">", "value": 50, "format": red_fill})
    data.conditional_format("E2:E21", {"type": "3_color_scale"})
    data.conditional_format("D2:D21", {"type": "data_bar"})
    data.set_column("A:E", 12)

    # Notes: image + autofilter with hidden (filtered-out) rows
    notes.write("A1", "Image below (24x16 PNG):")
    notes.insert_image("A2", "xcell.png", {"image_data": io.BytesIO(tiny_png())})
    notes.write_row("A10", ["Item", "Qty", "Status"], bold)
    items = [(f"item{i}", i, "open" if i % 2 else "closed") for i in range(1, 11)]
    for k, it in enumerate(items):
        notes.write_row(10 + k, 0, it)
    notes.autofilter("A10:C20")
    notes.filter_column("C", "Status == open")
    for k, it in enumerate(items):
        if it[2] != "open":
            notes.set_row(10 + k, options={"hidden": True})
    wb.close()
    normalize_zip(path)


# ---------------------------------------------------------------------------
# 4. macro.xlsm (xlsxwriter + real vbaProject.bin from calamine's vba.xlsm)
# ---------------------------------------------------------------------------

def gen_macro():
    import xlsxwriter

    with zipfile.ZipFile(os.path.join(VENDOR, "calamine_vba.xlsm")) as z:
        vba = z.read("xl/vbaProject.bin")
    path = out("macro.xlsm")
    wb = xlsxwriter.Workbook(path)
    wb.set_properties({"author": "xcell-corpus", "created": FIXED_DT})
    ws = wb.add_worksheet("Macro Sheet")
    ws.write("A1", "This workbook contains a VBA project (module testVBA).")
    ws.write("A2", 42)
    ws.write_formula("A3", "=A2*2", None, 84)
    wb.add_vba_project(io.BytesIO(vba), is_stream=True)
    wb.close()
    normalize_zip(path)
    shutil.copyfile(os.path.join(VENDOR, "calamine_vba.xlsm"), out("macro_calamine.xlsm"))


# ---------------------------------------------------------------------------
# 5. legacy.xls, sample.ods, xlsb copies
# ---------------------------------------------------------------------------

def gen_xls():
    import xlwt

    wb = xlwt.Workbook(encoding="utf-8")
    s1 = wb.add_sheet("Numbers")
    s1.write(0, 0, "n")
    s1.write(0, 1, "half")
    s1.write(0, 2, "money")
    money = xlwt.easyxf(num_format_str="#,##0.00")
    for i in range(1, 11):
        s1.write(i, 0, i)
        s1.write(i, 1, i / 2)
        s1.write(i, 2, i * 1234.5, money)
    s1.write(11, 0, xlwt.Formula("SUM(A2:A11)"))

    s2 = wb.add_sheet("Text")
    bold = xlwt.easyxf("font: bold on")
    s2.write(0, 0, "Greeting", bold)
    for r, t in enumerate(["Hello", "Café", "naïve", "日本語", "00123"], start=1):
        s2.write(r, 0, t)

    s3 = wb.add_sheet("Dates")
    dfmt = xlwt.easyxf(num_format_str="yyyy-mm-dd")
    dtfmt = xlwt.easyxf(num_format_str="yyyy-mm-dd hh:mm")
    s3.write(0, 0, "date")
    s3.write(1, 0, dt.date(2024, 3, 5), dfmt)
    s3.write(2, 0, dt.date(1999, 12, 31), dfmt)
    s3.write(3, 0, dt.datetime(2024, 3, 5, 14, 30), dtfmt)
    wb.save(out("legacy.xls"))


def gen_ods():
    from odf.opendocument import OpenDocumentSpreadsheet
    from odf.table import Table, TableCell, TableRow
    from odf.text import P

    doc = OpenDocumentSpreadsheet()

    def cell(value=None, text=None, **kw):
        c = TableCell(**kw)
        if text is not None:
            c.addElement(P(text=text))
        return c

    t1 = Table(name="Sheet One")
    hdr = TableRow()
    for h in ["name", "qty", "price", "date"]:
        hdr.addElement(cell(text=h, valuetype="string"))
    t1.addElement(hdr)
    data = [("apple", 3, 1.25, "2024-03-05"), ("banana", 12, 0.5, "2024-03-06"),
            ("cherry", 7, 3.75, "2023-12-31")]
    for name, qty, price, d in data:
        tr = TableRow()
        tr.addElement(cell(text=name, valuetype="string"))
        tr.addElement(cell(text=str(qty), valuetype="float", value=qty))
        tr.addElement(cell(text=f"{price}", valuetype="float", value=price))
        tr.addElement(cell(text=d, valuetype="date", datevalue=d))
        t1.addElement(tr)
    tr = TableRow()
    tr.addElement(cell(text="total", valuetype="string"))
    tr.addElement(cell(text="22", valuetype="float", value=22, formula="of:=SUM([.B2:.B4])"))
    t1.addElement(tr)
    doc.spreadsheet.addElement(t1)

    t2 = Table(name="Second")
    tr = TableRow()
    tr.addElement(cell(text="Unicode ✓ ₹ é", valuetype="string"))
    tr.addElement(cell(text="TRUE", valuetype="boolean", booleanvalue="true"))
    tr.addElement(cell(text="50%", valuetype="percentage", value=0.5))
    t2.addElement(tr)
    doc.spreadsheet.addElement(t2)

    path = out("sample.ods")
    doc.save(path)
    normalize_zip(path)


def copy_xlsb():
    shutil.copyfile(os.path.join(VENDOR, "calamine_issues.xlsb"), out("sample.xlsb"))
    shutil.copyfile(os.path.join(VENDOR, "calamine_any_sheets.xlsb"), out("sheet_states.xlsb"))


# ---------------------------------------------------------------------------
# 6. CSV variants
# ---------------------------------------------------------------------------

def wbytes(name, data):
    with open(out(name), "wb") as f:
        f.write(data)


def gen_csvs():
    wbytes("comma.csv", (
        "id,name,amount,date,active\n"
        "1,Alice,1234.5,2024-01-15,TRUE\n"
        "2,Bob,-42,2024-02-29,FALSE\n"
        "3,Chloé,0.001,2023-12-31,TRUE\n"
        "4,Dmitri,1e3,2024-06-01,FALSE\n"
        "5,Eve,,2024-07-04,TRUE\n").encode("utf-8"))

    wbytes("semicolon.csv", (
        "id;name;amount;ratio\n"
        "1;Anna;1.234,56;1,5\n"
        "2;Bernd;-12,75;0,25\n"
        "3;Claire;1000;2\n"
        '4;"Müller; Hans";3,14;0,5\n').encode("utf-8"))

    tsv = ("id\tname\tcity\tscore\n"
           "1\tSmith, John\tNew York\t88.5\n"
           "2\tO'Brien\tDublin\t92\n"
           "3\t\tBerlin\t\n"
           "4\tLi Wei\t北京\t75.25\n").encode("utf-8")
    wbytes("tabs.tsv", tsv)
    wbytes("tabs.txt", tsv)

    wbytes("quoted.csv", (
        "id,text,note\n"
        '1,"Hello, world",comma in quotes\n'
        '2,"She said ""hi""",doubled quotes\n'
        '3,"line one\nline two",embedded LF\n'
        '4,"",empty quoted\n'
        '5,"  spaces  ",leading/trailing spaces kept\n'
        '6,"multi\nline, with ""quotes"" and comma",everything\n'
        '7,"crlf\r\ninside",embedded CRLF\n').encode("utf-8"))

    wbytes("bom_crlf.csv", b"\xef\xbb\xbf" + (
        "name,city,amount\r\n"
        "Zoë,Zürich,12.5\r\n"
        "山田,東京,1000\r\n"
        "Ravi,मुंबई,₹500\r\n").encode("utf-8"))

    wbytes("latin1.csv", (
        "name,price,note\r\n"
        "Café,£3.50,résumé\r\n"
        "Jürgen,€12,naïve\r\n"
        "Ñoño,¥100,“smart quotes”\r\n").encode("cp1252"))

    wbytes("leading_zeros.csv", (
        "id,phone,price_text,sci_text,date,amount_inr,zip\n"
        '00123,+91 98765 43210,1.50,1e5,2024-03-05,"₹1,23,456.78",01234\n'
        '00007,09876543210,"1.50",1E-3,05/03/2024,"₹12,345.00",00501\n'
        '000,(555) 010-0199,0.10,2e10,3/5/2024,"₹0.50",90210\n'
        "\n"
        '42,555-0100,100,1.5e+3,05-Mar-2024,"₹1,00,00,000",07030\n'
        'short,row\n'
        '7,8,9,10,"March 5, 2024",₹99,11111,EXTRA1,EXTRA2\n'
        "\n"
        "0001,,,,20240305,,\n").encode("utf-8"))

    wbytes("no_trailing_newline.csv", b"a,b,c\n1,2,3\n4,5,6")


# ---------------------------------------------------------------------------
# 7. Large files
# ---------------------------------------------------------------------------

WORDS = ["alpha", "beta", "gamma", "delta", "epsilon", "zeta", "eta", "theta", "iota", "kappa",
         "lambda", "mu", "nu", "xi", "omicron", "pi", "rho", "sigma", "tau", "upsilon",
         "phi", "chi", "psi", "omega", "red", "green", "blue", "cyan", "magenta", "yellow",
         "north", "south", "east", "west", "open", "closed", "yes", "no", "N/A", "pending"]
LARGE_ROWS, LARGE_COLS = 10_000, 1_000  # total rows incl. header


def large_rows():
    """Yields data rows (without header) of python values; deterministic."""
    rng = random.Random(20240101)
    rnd, rint, uni, words = rng.random, rng.randint, rng.uniform, WORDS
    for _ in range(LARGE_ROWS - 1):
        row = []
        ap = row.append
        for _ in range(LARGE_COLS):
            if rnd() < 0.7:
                if rnd() < 0.5:
                    ap(rint(0, 99999))
                else:
                    ap(round(uni(-1000, 1000), 2))
            else:
                ap(words[rint(0, len(words) - 1)])
        yield row


LARGE_HEADER = [f"col_{i + 1}" for i in range(LARGE_COLS)]


def gen_large_xlsx_sst():
    """Hand-streamed xlsx with a shared-strings table (the way Excel stores strings)."""
    path = out("large_10k_x_1k.xlsx")
    letters = [col_letter(i) for i in range(LARGE_COLS)]
    sst = {}
    for w in LARGE_HEADER + WORDS:
        sst.setdefault(w, len(sst))
    str_count = 0
    last = f"{letters[-1]}{LARGE_ROWS}"

    def zi(name):
        z = zipfile.ZipInfo(name, date_time=FIXED_ZIP_DT)
        z.compress_type = zipfile.ZIP_DEFLATED
        z.external_attr = 0o644 << 16
        z.create_system = 0
        return z

    with zipfile.ZipFile(path, "w", compresslevel=6) as zf:
        zf.writestr(zi("[Content_Types].xml"), (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
            '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
            '<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>'
            '<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>'
            '</Types>'))
        zf.writestr(zi("_rels/.rels"), (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>'
            '</Relationships>'))
        zf.writestr(zi("xl/workbook.xml"), (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            '<sheets><sheet name="Data" sheetId="1" r:id="rId1"/></sheets></workbook>'))
        zf.writestr(zi("xl/_rels/workbook.xml.rels"), (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>'
            '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>'
            '<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml"/>'
            '</Relationships>'))
        zf.writestr(zi("xl/styles.xml"), (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            '<fonts count="2"><font><sz val="11"/><name val="Calibri"/></font>'
            '<font><b/><sz val="11"/><color rgb="FFFFFFFF"/><name val="Calibri"/></font></fonts>'
            '<fills count="3"><fill><patternFill patternType="none"/></fill>'
            '<fill><patternFill patternType="gray125"/></fill>'
            '<fill><patternFill patternType="solid"><fgColor rgb="FF305496"/><bgColor indexed="64"/></patternFill></fill></fills>'
            '<borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>'
            '<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>'
            '<cellXfs count="2"><xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>'
            '<xf numFmtId="0" fontId="1" fillId="2" borderId="0" xfId="0" applyFont="1" applyFill="1"/></cellXfs>'
            '<cellStyles count="1"><cellStyle name="Normal" xfId="0" builtinId="0"/></cellStyles>'
            '</styleSheet>'))

        with zf.open(zi("xl/worksheets/sheet1.xml"), "w", force_zip64=True) as fh:
            w = fh.write
            w(('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
               '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
               'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
               f'<dimension ref="A1:{last}"/>'
               '<sheetViews><sheetView workbookViewId="0"><pane ySplit="1" topLeftCell="A2" '
               'activePane="bottomLeft" state="frozen"/><selection pane="bottomLeft" activeCell="A2" sqref="A2"/>'
               '</sheetView></sheetViews><sheetFormatPr defaultRowHeight="15"/><sheetData>').encode())
            parts = [f'<row r="1" spans="1:{LARGE_COLS}">']
            for i, h in enumerate(LARGE_HEADER):
                s = ' s="1"' if i < 5 else ""
                parts.append(f'<c r="{letters[i]}1"{s} t="s"><v>{sst[h]}</v></c>')
                str_count += 1
            parts.append("</row>")
            w("".join(parts).encode())
            for r, row in enumerate(large_rows(), start=2):
                rs = str(r)
                parts = [f'<row r="{rs}" spans="1:{LARGE_COLS}">']
                ap = parts.append
                for i, v in enumerate(row):
                    if type(v) is str:
                        ap(f'<c r="{letters[i]}{rs}" t="s"><v>{sst[v]}</v></c>')
                        str_count += 1
                    else:
                        ap(f'<c r="{letters[i]}{rs}"><v>{v!r}</v></c>')
                ap("</row>")
                w("".join(parts).encode())
            w(b"</sheetData></worksheet>")

        items = "".join(f"<si><t>{s}</t></si>" for s in sst)
        zf.writestr(zi("xl/sharedStrings.xml"), (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            f'count="{str_count}" uniqueCount="{len(sst)}">{items}</sst>'))


def gen_large_xlsx_inline():
    """Same data via xlsxwriter constant_memory (strings become inlineStr)."""
    import xlsxwriter

    path = out("large_10k_x_1k_inline.xlsx")
    wb = xlsxwriter.Workbook(path, {"constant_memory": True})
    wb.set_properties({"author": "xcell-corpus", "created": FIXED_DT})
    ws = wb.add_worksheet("Data")
    hdr = wb.add_format({"bold": True, "font_color": "#FFFFFF", "bg_color": "#305496"})
    for i, h in enumerate(LARGE_HEADER):
        ws.write_string(0, i, h, hdr if i < 5 else None)
    ws.freeze_panes(1, 0)
    wn, ws_ = ws.write_number, ws.write_string
    for r, row in enumerate(large_rows(), start=1):
        for c, v in enumerate(row):
            if type(v) is str:
                ws_(r, c, v)
            else:
                wn(r, c, v)
    wb.close()
    normalize_zip(path)


def gen_large_csv():
    with open(out("large_10k_x_1k.csv"), "w", encoding="utf-8", newline="") as f:
        f.write(",".join(LARGE_HEADER) + "\n")
        for row in large_rows():
            f.write(",".join(v if type(v) is str else repr(v) for v in row) + "\n")


def gen_tall_csv():
    rng = random.Random(99)
    base = dt.date(2020, 1, 1)
    hdr = ["id", "date", "region", "product", "qty", "price", "amount", "flag", "code", "note",
           "m1", "m2", "m3", "m4", "m5", "m6", "m7", "m8", "m9", "m10"]
    with open(out("large_200k_x_20.csv"), "w", encoding="utf-8", newline="") as f:
        f.write(",".join(hdr) + "\n")
        for i in range(1, 200_000):  # 199,999 data rows + header = 200,000 lines
            qty = rng.randint(1, 500)
            price = round(rng.uniform(0.5, 999.99), 2)
            d = base + dt.timedelta(days=rng.randint(0, 1825))
            fields = [str(i), d.isoformat(), REGIONS[rng.randint(0, 3)], PRODUCTS[rng.randint(0, 4)],
                      str(qty), repr(price), repr(round(qty * price, 2)),
                      "TRUE" if rng.random() < 0.5 else "FALSE", f"{rng.randint(0, 99999):05d}",
                      WORDS[rng.randint(0, len(WORDS) - 1)]]
            fields += [repr(round(rng.gauss(0, 100), 3)) for _ in range(10)]
            f.write(",".join(fields) + "\n")


# ---------------------------------------------------------------------------

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--skip-large", action="store_true")
    ap.add_argument("--only", help="comma-separated step names")
    args = ap.parse_args()
    os.makedirs(OUT, exist_ok=True)
    steps = [("styled", gen_styled), ("formulas", gen_formulas), ("rich", gen_rich),
             ("macro", gen_macro), ("xls", gen_xls), ("ods", gen_ods), ("xlsb", copy_xlsb),
             ("csv", gen_csvs)]
    if not args.skip_large:
        steps += [("large_xlsx", gen_large_xlsx_sst), ("large_xlsx_inline", gen_large_xlsx_inline),
                  ("large_csv", gen_large_csv), ("tall_csv", gen_tall_csv)]
    if args.only:
        wanted = set(args.only.split(","))
        steps = [s for s in steps if s[0] in wanted]
    for name, fn in steps:
        t = time.perf_counter()
        fn()
        print(f"{name:20s} {time.perf_counter() - t:7.2f}s", flush=True)


if __name__ == "__main__":
    sys.exit(main())
