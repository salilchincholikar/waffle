#!/usr/bin/env python3
"""Re-open every generated file with independent readers and assert the key
facts documented in testdata/README.md. Run after generate.py."""
import datetime as dt
import os
import sys
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "..", "..", "testdata", "files")


def p(name):
    return os.path.join(OUT, name)


def check_styled():
    from openpyxl import load_workbook
    wb = load_workbook(p("styled.xlsx"))
    assert wb.sheetnames == ["Styled"]
    ws = wb["Styled"]
    assert ws.dimensions == "A1:H38", ws.dimensions
    assert ws.freeze_panes == "B2"
    assert sorted(map(str, ws.merged_cells.ranges)) == ["D2:F2", "D4:D7", "D9:F11"]
    assert ws.row_dimensions[37].hidden and ws.column_dimensions["G"].hidden
    assert ws["B2"].font.b and ws["B3"].font.i and ws["B4"].font.u == "single" and ws["B5"].font.strike
    assert ws["B6"].font.name == "Georgia" and ws["B6"].font.sz == 18
    assert ws["B8"].font.color.rgb == "FFFF0000"
    assert ws["B9"].font.color.theme == 4 and ws["B10"].font.color.indexed == 12
    assert ws["B11"].fill.fgColor.rgb == "FFFFFF00"
    assert ws["B15"].border.left.style == "dashed" and ws["B15"].border.left.color.rgb == "FFFF0000"
    assert ws["B19"].alignment.indent == 2 and ws["B18"].alignment.wrap_text
    assert ws["B21"].number_format == "#,##0.00" and ws["B21"].value == 1234567.891
    assert ws["B24"].value == dt.datetime(2024, 3, 5) and ws["B24"].number_format == "m/d/yyyy"
    assert ws["B30"].number_format == "[$₹-4009] #,##0.00"
    assert ws["B34"].value == "00123" and ws["B34"].number_format == "@"
    assert ws.column_dimensions["B"].width == 26 and ws.row_dimensions[16].height == 40


def check_formulas():
    from openpyxl import load_workbook
    wb = load_workbook(p("formulas.xlsx"))
    assert wb.sheetnames == ["Calc", "Other Sheet"]
    ws = wb["Calc"]
    assert ws["B2"].value == "=A2*2" and ws["B7"].value == "=A7*2"     # shared child expanded
    assert ws["C11"].value == "=A11*$F$1"
    assert ws["C13"].value == "=SUM(C2:C11)"
    assert ws["E2"].value == "='Other Sheet'!A1"
    assert ws["H2"].value == "=SUM(A:A)"
    assert "TaxRate" in wb.defined_names
    vals = load_workbook(p("formulas.xlsx"), data_only=True)["Calc"]
    assert [vals[f"B{r}"].value for r in range(2, 12)] == list(range(2, 21, 2))
    assert vals["C13"].value == 165 and vals["H2"].value == 55 and vals["H4"].value == 18
    assert vals["D2"].value == "small" and vals["D11"].value == "big"
    assert vals["J6"].value == 770 and vals["J3"].value == 8
    xml = zipfile.ZipFile(p("formulas.xlsx")).read("xl/worksheets/sheet1.xml").decode()
    assert xml.count('t="shared"') == 22 and '<f t="shared" ref="B2:B11" si="0">A2*2</f>' in xml
    assert '<f t="shared" ref="B13:C13" si="2">SUM(B2:B11)</f>' in xml
    assert '<f t="array" ref="J2:J4">A2:A4*B2:B4</f>' in xml


def check_rich():
    from openpyxl import load_workbook
    wb = load_workbook(p("rich.xlsx"), rich_text=True)
    assert wb.sheetnames == ["Summary", "Data", "Hidden Calc", "Notes"]
    assert wb["Hidden Calc"].sheet_state == "hidden"
    s = wb["Summary"]
    assert str(s["A1"].value) == "Bold red plain italic blue" and len(s["A1"].value) == 3
    assert s["A6"].hyperlink.target == "https://example.com/xcell"
    assert s["B3"].comment.text == "Pick a region from the list"
    assert len(s.data_validations.dataValidation) == 2
    d = wb["Data"]
    assert list(d.tables) == ["SalesTable"] and d.tables["SalesTable"].ref == "A1:E21"
    assert len(d.conditional_formatting) == 3
    assert wb["Notes"].auto_filter.ref == "A10:C20"
    z = zipfile.ZipFile(p("rich.xlsx"))
    names = z.namelist()
    for part in ["xl/charts/chart1.xml", "xl/media/image1.png", "xl/comments1.xml", "xl/tables/table1.xml"]:
        assert part in names, part


def check_macro():
    from openpyxl import load_workbook
    for f in ["macro.xlsm", "macro_calamine.xlsm"]:
        wb = load_workbook(p(f), keep_vba=True)
        assert wb.vba_archive is not None
        z = zipfile.ZipFile(p(f))
        assert "xl/vbaProject.bin" in z.namelist()
        assert b"macroEnabled.main+xml" in z.read("[Content_Types].xml")
    wb = load_workbook(p("macro.xlsm"), keep_vba=True)
    assert wb.sheetnames == ["Macro Sheet"] and wb["Macro Sheet"]["A2"].value == 42


def check_xls():
    import xlrd
    b = xlrd.open_workbook(p("legacy.xls"))
    assert b.sheet_names() == ["Numbers", "Text", "Dates"]
    n = b.sheet_by_name("Numbers")
    assert (n.nrows, n.ncols) == (12, 3) and n.cell_value(10, 2) == 12345.0
    t = b.sheet_by_name("Text")
    assert t.cell_value(4, 0) == "日本語"
    d = b.sheet_by_name("Dates")
    assert xlrd.xldate_as_datetime(d.cell_value(1, 0), b.datemode) == dt.datetime(2024, 3, 5)


def check_ods():
    from odf.opendocument import load
    from odf.table import Table
    doc = load(p("sample.ods"))
    names = [t.getAttribute("name") for t in doc.spreadsheet.getElementsByType(Table)]
    assert names == ["Sheet One", "Second"], names


def check_xlsb():
    from pyxlsb import open_workbook
    with open_workbook(p("sample.xlsb")) as wb:
        assert wb.sheets == ["datatypes", "issue2", "Sheet1", "issue5", "issue6", "spc_chrs"]
        with wb.get_sheet("issue2") as sh:
            assert [[c.v for c in r] for r in sh.rows()] == [[1.0, "a"], [2.0, "b"], [3.0, "c"]]
    with open_workbook(p("sheet_states.xlsb")) as wb:
        assert wb.sheets == ["Visible", "Hidden", "VeryHidden", "Chart"]


def check_csv():
    raw = open(p("latin1.csv"), "rb").read()
    assert b"\x80" in raw and b"\xa3" in raw  # € and £ in cp1252
    assert open(p("bom_crlf.csv"), "rb").read().startswith(b"\xef\xbb\xbf")
    assert not open(p("no_trailing_newline.csv"), "rb").read().endswith(b"\n")
    import csv
    rows = list(csv.reader(open(p("quoted.csv"), newline="", encoding="utf-8")))
    assert len(rows) == 8 and rows[3][1] == "line one\nline two" and rows[2][1] == 'She said "hi"'


def check_large():
    if not os.path.exists(p("large_10k_x_1k.xlsx")):
        print("  (large files absent, skipped)")
        return
    from openpyxl import load_workbook
    for f in ["large_10k_x_1k.xlsx", "large_10k_x_1k_inline.xlsx"]:
        wb = load_workbook(p(f), read_only=True)
        ws = wb["Data"]
        assert ws.calculate_dimension() == "A1:ALL10000", ws.calculate_dimension()
        first = next(ws.iter_rows(min_row=1, max_row=2, values_only=True))
        assert first[0] == "col_1" and first[999] == "col_1000"
        last = next(ws.iter_rows(min_row=10000, max_row=10000, values_only=True))
        assert len(last) == 1000
        wb.close()
    with open(p("large_200k_x_20.csv"), "rb") as fh:
        assert sum(1 for _ in fh) == 200_000


def main():
    for name, fn in list(globals().items()):
        if name.startswith("check_"):
            fn()
            print("ok", name)


if __name__ == "__main__":
    sys.exit(main())
